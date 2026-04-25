"""
extractor.py — Handles text extraction from .txt, .docx, .pdf, and .xlsx files.
Normalizes whitespace and encoding for clean comparison.
"""

import re
import io
import chardet


def normalize_text(text: str) -> str:
    """
    Normalize extracted text:
    - Collapse multiple spaces/tabs into single space
    - Normalize line endings
    - Strip trailing whitespace per line
    - Remove blank-only lines (keep intentional paragraph breaks as single blank line)
    """
    # Normalize line endings
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    lines = text.split("\n")
    normalized = []
    prev_blank = False

    for line in lines:
        # Collapse internal whitespace, strip edges
        line = re.sub(r"[ \t]+", " ", line).strip()
        if line == "":
            if not prev_blank:
                normalized.append("")
            prev_blank = True
        else:
            normalized.append(line)
            prev_blank = False

    # Strip leading/trailing blank lines
    while normalized and normalized[0] == "":
        normalized.pop(0)
    while normalized and normalized[-1] == "":
        normalized.pop()

    return "\n".join(normalized)


def extract_txt(file_bytes: bytes) -> str:
    """Extract text from a plain .txt file, auto-detecting encoding."""
    detected = chardet.detect(file_bytes)
    encoding = detected.get("encoding") or "utf-8"
    try:
        text = file_bytes.decode(encoding, errors="replace")
    except Exception:
        text = file_bytes.decode("utf-8", errors="replace")
    return normalize_text(text)


def extract_docx(file_bytes: bytes) -> str:
    """Extract text from a .docx file preserving paragraph structure."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
            else:
                paragraphs.append("")  # preserve paragraph breaks

        # Also pull text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(f"[TABLE ROW] {row_text}")

        return normalize_text("\n".join(paragraphs))
    except ImportError:
        raise ImportError("python-docx is required: pip install python-docx")
    except Exception as e:
        raise ValueError(f"Failed to read .docx: {e}")


def extract_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF using pdfplumber (preferred) with PyPDF2 fallback."""
    # Try pdfplumber first (better text extraction)
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = []
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if text:
                    pages.append(f"[PAGE {i+1}]\n{text}")
            if pages:
                return normalize_text("\n\n".join(pages))
    except ImportError:
        pass
    except Exception:
        pass

    # Fallback: PyPDF2
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        pages = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[PAGE {i+1}]\n{text}")
        if pages:
            return normalize_text("\n\n".join(pages))
    except ImportError:
        pass
    except Exception:
        pass

    raise ValueError(
        "Could not extract PDF text. Install pdfplumber: pip install pdfplumber"
    )


def extract_xlsx(file_bytes: bytes) -> str:
    """Extract text from an .xlsx file, sheet by sheet, row by row."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"[SHEET: {sheet_name}]")
            for row in ws.iter_rows():
                cells = []
                for cell in row:
                    val = cell.value
                    if val is not None:
                        cells.append(str(val).strip())
                if cells:
                    lines.append(" | ".join(cells))
            lines.append("")  # blank line between sheets
        return normalize_text("\n".join(lines))
    except ImportError:
        raise ImportError("openpyxl is required: pip install openpyxl")
    except Exception as e:
        raise ValueError(f"Failed to read .xlsx: {e}")


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Route file to the correct extractor based on extension.
    Returns normalized text string.
    """
    ext = filename.rsplit(".", 1)[-1].lower()

    if ext == "txt":
        return extract_txt(file_bytes)
    elif ext == "docx":
        return extract_docx(file_bytes)
    elif ext == "pdf":
        return extract_pdf(file_bytes)
    elif ext == "xlsx":
        return extract_xlsx(file_bytes)
    else:
        raise ValueError(
            f"Unsupported file type: .{ext}\n"
            "Supported types: .txt, .docx, .pdf, .xlsx"
        )
